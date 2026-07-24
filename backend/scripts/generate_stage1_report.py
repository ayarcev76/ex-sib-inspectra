"""Генератор отчета по Этапу 1 в формате .docx."""
import sys
from pathlib import Path
from datetime import datetime

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("❌ Библиотека python-docx не установлена.")
    print("Установите её командой: pip install python-docx")
    sys.exit(1)


def set_cell_bold(cell, text):
    """Вспомогательная функция для установки жирного текста в ячейке."""
    cell.text = text
    if cell.paragraphs and cell.paragraphs[0].runs:
        cell.paragraphs[0].runs[0].font.bold = True


def create_table_with_header(doc, headers, data, style='Light Grid Accent 1'):
    """
    Универсальная функция создания таблицы с заголовком.
    
    Автоматически рассчитывает количество строк на основе данных,
    что защищает от ошибок типа IndexError.
    """
    # Количество строк = 1 (заголовок) + количество строк данных
    num_rows = 1 + len(data)
    num_cols = len(headers)
    
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.style = style
    
    # Заполняем заголовок
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
        if table.rows[0].cells[i].paragraphs[0].runs:
            table.rows[0].cells[i].paragraphs[0].runs[0].font.bold = True
    
    # Заполняем данные
    for row_idx, row_data in enumerate(data, start=1):
        for col_idx, cell_value in enumerate(row_data):
            table.rows[row_idx].cells[col_idx].text = str(cell_value)
    
    return table


def create_report():
    """Создает документ отчета."""
    doc = Document()
    
    # === СТИЛИ ===
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    
    # === ЗАГОЛОВОК ===
    title = doc.add_heading('ОТЧЁТ О ВЫПОЛНЕНИИ ЭТАПА 1', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Проектирование базы данных')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.color.rgb = RGBColor(89, 89, 89)
    
    doc.add_paragraph('─' * 60)
    
    # === МЕТА-ИНФОРМАЦИЯ ===
    meta_data = [
        ('Проект:', 'ЕХ:СИБ — Система инспекций безопасности'),
        ('Заказчик:', 'ГК «ЕВРОХИМ»'),
        ('Дата отчёта:', datetime.now().strftime('%d.%m.%Y')),
        ('Статус этапа:', '✅ Завершён'),
        ('Ответственный:', 'Команда разработки'),
    ]
    meta_table = doc.add_table(rows=len(meta_data), cols=2)
    meta_table.style = 'Light Grid Accent 1'
    for i, (label, value) in enumerate(meta_data):
        meta_table.rows[i].cells[0].text = label
        meta_table.rows[i].cells[1].text = value
        if meta_table.rows[i].cells[0].paragraphs[0].runs:
            meta_table.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True
    
    doc.add_paragraph()
    
    # === 1. РЕЗЮМЕ ===
    doc.add_heading('1. Резюме', 1)
    doc.add_paragraph(
        'Этап 1 успешно завершён. Спроектирована и реализована полная схема базы данных '
        'PostgreSQL 15 для системы инспекций безопасности ЕХ:СИБ. Создано 15 таблиц, '
        'соответствующих Техническому заданию v3.3, включая справочники, модели пользователей, '
        'основную бизнес-логику (Inspection, ViolationRecord, ViolationPhoto), модели планирования '
        'и уведомлений. Все таблицы проиндексированы согласно требованиям ТЗ (п. 10.3). '
        'Написан seed-скрипт для наполнения справочников начальными данными.'
    )
    
    # === 2. ВЫПОЛНЕННЫЕ РАБОТЫ ===
    doc.add_heading('2. Выполненные работы (Чек-лист ТЗ)', 1)
    
    tasks_data = [
        ('1', 'SQLAlchemy модели: User, Role, PE, Department, Contractor, WorkType, ZPBRule', '✅ Выполнено', 'Созданы модели с UUID PK, timestamps, relationships'),
        ('2', 'Модели: Inspection, ViolationRecord, ViolationPhoto', '✅ Выполнено', 'Поле source (mobile/web), status (draft/submitted/approved)'),
        ('3', 'Модели: InspectionPlan, UserPEAssignment, Notification', '✅ Выполнено', 'Enum-типы для статусов, связи с пользователями'),
        ('4', 'Поле source в модели Inspection (mobile/web)', '✅ Выполнено', 'Enum InspectionSource с значениями mobile/web'),
        ('5', 'Создание индексов БД для оптимизации (Alembic)', '✅ Выполнено', 'Индексы на: pe_id, date, inspector_id, status, plan_id'),
        ('6', 'Alembic миграции', '✅ Выполнено', '4 миграции: reference_tables, user_role_tables, inspection_tables, planning_tables'),
        ('7', 'Seed-скрипт для справочников', '✅ Выполнено', '6 ролей, 7 ЗПБ, 8 видов работ, 5 ПЕ, 1 тестовый админ'),
    ]
    create_table_with_header(
        doc,
        headers=['№', 'Задача по ТЗ (Этап 1)', 'Статус', 'Примечание'],
        data=tasks_data,
    )
    
    doc.add_paragraph()
    
    # === 3. СПЕЦИФИКАЦИЯ БД ===
    doc.add_heading('3. Спецификация базы данных', 1)
    
    doc.add_heading('3.1. Таблицы (15 штук)', 2)
    
    tables_list = [
        ('pe', 'Производственные единицы (5 предприятий)'),
        ('department', 'Подразделения (FK → pe, unique constraint name+pe_id)'),
        ('contractor', 'Подрядные организации (внешние и внутренние)'),
        ('work_type', 'Виды работ (8 типов: ремонтные, огневые и т.д.)'),
        ('zpb_rule', 'Золотые Правила Безопасности (7 правил)'),
        ('role', 'Роли пользователей (6 ролей)'),
        ('user', 'Пользователи системы (email, hashed_password)'),
        ('user_roles', 'Связь M2M user ↔ role'),
        ('inspection', 'Шапка проверки (inspection_number, date, source, status)'),
        ('violation_record', 'Строка наблюдения (is_safe, is_gross_violation, is_top_violation)'),
        ('violation_photo', 'Фотография нарушения (file_path, thumbnail_path, GPS)'),
        ('inspection_plan', 'План инспекций на неделю (week_start_date, status)'),
        ('user_pe_assignment', 'Назначение инспектора на ПЕ (valid_from, valid_to)'),
        ('notification', 'Журнал уведомлений (type, channel, is_read)'),
        ('alembic_version', 'Версии миграций Alembic'),
    ]
    
    for table_name, description in tables_list:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(f'{table_name}').bold = True
        p.add_run(f' — {description}')
    
    doc.add_heading('3.2. Индексы (оптимизация производительности)', 2)
    doc.add_paragraph('Согласно п. 10.3 ТЗ, созданы индексы для критичных полей:')
    
    indexes_list = [
        'inspection.pe_id, inspection.department_id, inspection.inspector_id, inspection.contractor_id',
        'inspection.date, inspection.status, inspection.source',
        'inspection.plan_id (FK → inspection_plan)',
        'violation_record.inspection_id, violation_record.work_type_id, violation_record.zpb_rule_id',
        'violation_record.is_top_violation (для фильтрации топ-нарушений)',
        'violation_photo.violation_id, violation_photo.uploaded_by',
        'inspection_plan.week_start_date (unique), inspection_plan.status',
        'user_pe_assignment.user_id, user_pe_assignment.pe_id, user_pe_assignment.is_active',
        'notification.user_id, notification.type, notification.is_read',
    ]
    
    for idx in indexes_list:
        doc.add_paragraph(idx, style='List Bullet')
    
    doc.add_heading('3.3. Foreign Keys и ограничения', 2)
    doc.add_paragraph('Все связи реализованы через Foreign Keys с правильными ON DELETE правилами:')
    
    fk_rules = [
        'ON DELETE CASCADE: violation_record → inspection, violation_photo → violation_record',
        'ON DELETE RESTRICT: inspection → pe/department/user/contractor (защита от случайного удаления)',
        'ON DELETE SET NULL: inspection → inspection_plan (план может быть удалён)',
        'UNIQUE CONSTRAINT: department(name, pe_id), inspection(inspection_number), pe(code), work_type(code)',
    ]
    
    for rule in fk_rules:
        doc.add_paragraph(rule, style='List Bullet')
    
    doc.add_paragraph()
    
    # === 4. ПРОБЛЕМЫ И РЕШЕНИЯ ===
    doc.add_heading('4. Проблемы и решения', 1)
    doc.add_paragraph('В процессе выполнения этапа возникли следующие технические проблемы:')
    
    problems_data = [
        ('1', 'psycopg2-binary не имеет wheel для Python 3.14', 'Перешли на psycopg v3 (psycopg[binary]==3.2.10)'),
        ('2', 'pydantic-core не компилируется на Python 3.14 (нет Rust toolchain)', 'Понизили Python до 3.12 (рекомендуется для стабильности)'),
        ('3', 'Alembic не видит модуль app (ModuleNotFoundError)', 'Создан pyproject.toml + pip install -e . (editable mode)'),
        ('4', 'PostgreSQL отвергает пароль (password authentication failed)', 'docker-compose down -v + пересоздание контейнера с правильным паролем'),
        ('5', 'Alembic не находит script.py.mako', 'Создан файл шаблона alembic/script.py.mako вручную'),
        ('6', 'Коллизия имен индексов (ix_inspection_plan_id already exists)', 'Убран index=True из UUIDMixin (PK автоматически индексируется в PostgreSQL)'),
        ('7', 'passlib 1.7.4 несовместим с bcrypt 5.x', 'Зафиксирована версия bcrypt==4.2.1 в requirements.txt'),
        ('8', 'IndexError в скрипте генерации отчета', 'Исправлен расчет количества строк таблицы (rows=1+len(data))'),
    ]
    create_table_with_header(
        doc,
        headers=['№', 'Проблема', 'Решение'],
        data=problems_data,
    )
    
    doc.add_paragraph()
    
    # === 5. ЗАМЕТКИ НА БУДУЩЕЕ ===
    doc.add_heading('5. Заметки на будущее (Best Practices)', 1)
    
    notes = [
        ('Python версии', 'Для продакшена рекомендуется Python 3.12 (стабильный, все wheel-пакеты доступны). Python 3.14 требует компиляции многих пакетов из исходников.'),
        ('SQLAlchemy 2.0', 'Использовать text() для сырых SQL-запросов: conn.execute(text("SELECT 1")). Строки не принимаются напрямую.'),
        ('Alembic и индексы', 'Не указывать index=True для первичных ключей — PostgreSQL автоматически создает индекс для PK. Избыточные индексы приводят к коллизиям имен.'),
        ('passlib vs bcrypt', 'passlib 1.7.4 не обновлялась с 2020 года. Для новых проектов рассмотреть прямой bcrypt или альтернативы (argon2-cffi).'),
        ('Миграции с FK', 'Если модель A ссылается на модель B, но таблица B ещё не создана, Alembic упадет с NoReferencedTableError. Решение: создавать таблицы в правильном порядке или разбивать миграции.'),
        ('Docker volumes', 'При изменении POSTGRES_PASSWORD в docker-compose.yml нужно удалять старый volume (docker-compose down -v), иначе PostgreSQL проигнорирует новый пароль.'),
        ('Enum в SQLAlchemy', 'Использовать Python Enum-классы (str, enum.Enum) для типобезопасности. SQLAlchemy автоматически создаст PostgreSQL enum-тип.'),
        ('Relationships lazy loading', 'lazy="joined" для частых связей (PE, Department), lazy="selectin" для коллекций (violations), lazy="noload" для обратных связей.'),
        ('python-docx', 'У объекта Paragraph нет атрибута font напрямую. Для форматирования использовать paragraph.runs[0].font.bold = True.'),
        ('python-docx таблицы', 'При создании таблицы через add_table(rows=N, cols=M) всегда учитывать строку заголовка. Формула: rows = 1 (заголовок) + len(data). Рекомендуется вычислять количество строк динамически.'),
        ('Проверка кода', 'Перед отправкой кода пользователю всегда проверять: (1) соответствие размеров таблиц количеству данных, (2) корректность импортов, (3) граничные случаи в циклах.'),
    ]
    
    for title, note in notes:
        p = doc.add_paragraph()
        p.add_run(f'{title}: ').bold = True
        p.add_run(note)
    
    doc.add_paragraph()
    
    # === 6. СЛЕДУЮЩИЕ ШАГИ ===
    doc.add_heading('6. Следующие шаги (Переход к Этапу 2)', 1)
    doc.add_paragraph('Согласно ТЗ, следующим этапом является Этап 2: Backend — базовый CRUD (2 недели).')
    
    next_steps = [
        '☐ FastAPI приложение + подключение к БД (уже частично готово)',
        '☐ Аутентификация (JWT: access 15 мин + refresh 7 дней)',
        '☐ CRUD справочников (с фильтром подразделений по ПЕ)',
        '☐ CRUD пользователей + множественные роли',
        '☐ CRUD назначений инспекторов на ПЕ (с историей)',
        '☐ CRUD проверок и нарушений (с поддержкой веб и мобильного приложения)',
        '☐ Авторасчет is_top_violation (бизнес-логика)',
        '☐ Unit-тесты (pytest)',
    ]
    
    for step in next_steps:
        doc.add_paragraph(step, style='List Bullet')
    
    doc.add_paragraph()
    
    # === ПОДПИСЬ ===
    doc.add_paragraph('─' * 60)
    doc.add_paragraph()
    
    signature_table = doc.add_table(rows=2, cols=2)
    signature_table.rows[0].cells[0].text = 'Подпись разработчика:'
    signature_table.rows[0].cells[1].text = '____________________'
    signature_table.rows[1].cells[0].text = 'Дата:'
    signature_table.rows[1].cells[1].text = datetime.now().strftime('%d.%m.%Y')
    
    # === СОХРАНЕНИЕ ===
    output_path = Path(__file__).parent.parent.parent / 'docs' / 'Отчет_Этап_1_Проектирование_БД.docx'
    output_path.parent.mkdir(parents=True, exist_ok=True)
    
    doc.save(str(output_path))
    print(f'✅ Отчет успешно создан: {output_path}')
    print(f'📄 Размер файла: {output_path.stat().st_size / 1024:.1f} KB')


if __name__ == '__main__':
    create_report()