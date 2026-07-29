"""Сервис экспорта карты наблюдения в PDF и DOCX."""
import os
import uuid
import tempfile
import logging
from io import BytesIO
from pathlib import Path
from datetime import datetime
from sqlalchemy.orm import Session
from xhtml2pdf import pisa
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image as PILImage
from app.models.inspections import Inspection, ViolationRecord
from app.models.references import PE, Department, Contractor, WorkType, ZPBRule
from app.models.users import User
from app.core.minio_client import get_file_bytes

logger = logging.getLogger(__name__)

# ==================== КОНФИГУРАЦИЯ ШРИФТОВ ====================
# Используем бесплатный шрифт DejaVu Sans с поддержкой кириллицы
FONT_REGULAR_PATH = Path(__file__).parent.parent / "static" / "fonts" / "DejaVuSans.ttf"

_font_registered = False

def _register_fonts():
    """
    Регистрация шрифта в reportlab и xhtml2pdf.
    Шрифт DejaVu Sans поддерживает кириллицу и лежит в локальной папке проекта.
    """
    global _font_registered
    if _font_registered:
        return
    
    if not FONT_REGULAR_PATH.exists():
        raise RuntimeError(
            f"Шрифт DejaVu Sans не найден: {FONT_REGULAR_PATH}\n"
            f"Убедитесь, что файл DejaVuSans.ttf существует в папке backend/app/static/fonts/"
        )
    
    # Регистрируем шрифт в глобальном реестре reportlab
    pdfmetrics.registerFont(TTFont('DejaVuSans', str(FONT_REGULAR_PATH)))
    
    # Регистрируем шрифт в xhtml2pdf через addMapping
    from xhtml2pdf.config.html2pdfdata import addMapping
    font_uri = FONT_REGULAR_PATH.absolute().as_uri()
    addMapping('DejaVuSans', 0, 0, font_uri)  # 0, 0 - normal, normal (не bold, не italic)
    
    _font_registered = True
    logger.info(f"✅ Шрифт 'DejaVuSans' зарегистрирован: {FONT_REGULAR_PATH}")

# ==================== СБОР ДАННЫХ ====================
def _get_inspection_data(db: Session, inspection_id: uuid.UUID) -> dict:
    """Сбор всех данных проверки для экспорта."""
    logger.info(f"📊 Сбор данных для экспорта проверки {inspection_id}")
    
    inspection = db.query(Inspection).filter(Inspection.id == inspection_id).first()
    if not inspection:
        raise ValueError("Карта наблюдения не найдена")
    
    pe = db.query(PE).filter(PE.id == inspection.pe_id).first()
    dept = db.query(Department).filter(Department.id == inspection.department_id).first()
    inspector = db.query(User).filter(User.id == inspection.inspector_id).first()
    contractor = db.query(Contractor).filter(Contractor.id == inspection.contractor_id).first() if inspection.contractor_id else None
    
    violations = db.query(ViolationRecord).filter(
        ViolationRecord.inspection_id == inspection_id
    ).order_by(ViolationRecord.order).all()
    
    violations_data = []
    for v in violations:
        work_type = db.query(WorkType).filter(WorkType.id == v.work_type_id).first() if v.work_type_id else None
        zpb_rule = db.query(ZPBRule).filter(ZPBRule.id == v.zpb_rule_id).first() if v.zpb_rule_id else None
        
        photos_temp_paths = []
        for photo in v.photos:
            try:
                logger.info(f"📥 Загрузка фото {photo.id} из MinIO: {photo.file_path}")
                file_bytes = get_file_bytes(photo.file_path)
                logger.info(f"✅ Фото загружено: {len(file_bytes)} байт")
                
                img_stream = BytesIO(file_bytes)
                pil_image = PILImage.open(img_stream)
                if pil_image.mode in ('RGBA', 'P'):
                    pil_image = pil_image.convert('RGB')
                
                temp_file = tempfile.NamedTemporaryFile(
                    delete=False, suffix='.jpg', prefix='inspectra_'
                )
                pil_image.save(temp_file.name, format='JPEG', quality=85)
                temp_file.close()
                
                clean_path = temp_file.name.replace('\\', '/')
                logger.info(f"💾 Временный файл создан: {clean_path}")
                
                photos_temp_paths.append({
                    'path': temp_file.name,
                    'uri': clean_path,
                    'caption': photo.caption,
                    'gps': f"{photo.gps_latitude:.4f}, {photo.gps_longitude:.4f}" if photo.gps_latitude and photo.gps_longitude else None,
                })
            except Exception as e:
                logger.error(f"⚠️ Не удалось обработать фото {photo.id}: {e}", exc_info=True)
        
        violations_data.append({
            'order': v.order,
            'is_safe': v.is_safe,
            'description': v.violation_description,
            'work_type': work_type.name if work_type else None,
            'is_gross': v.is_gross_violation,
            'is_stopped': v.is_work_stopped,
            'zpb_rule': f"{zpb_rule.number}. {zpb_rule.name}" if zpb_rule else None,
            'is_top': v.is_top_violation,
            'photos': photos_temp_paths,
        })
    
    return {
        'inspection': inspection,
        'pe': pe.name if pe else '—',
        'department': dept.name if dept else '—',
        'inspector': inspector.full_name if inspector else '—',
        'contractor': contractor.name if contractor else '—',
        'violations': violations_data,
    }

def _cleanup_temp_files(violations_data: list):
    """Удаление временных файлов фотографий."""
    for v in violations_data:
        for photo in v.get('photos', []):
            try:
                if os.path.exists(photo['path']):
                    os.remove(photo['path'])
                    logger.info(f"🗑️ Временный файл удален: {photo['path']}")
            except Exception as e:
                logger.warning(f"Не удалось удалить временный файл {photo['path']}: {e}")

# ==================== PDF EXPORT ====================
def generate_inspection_pdf(inspection_id: uuid.UUID, db: Session) -> bytes:
    """Генерация PDF-отчёта по карте наблюдения."""
    data = None
    try:
        logger.info(f"🚀 Начало генерации PDF для проверки {inspection_id}")
        
        # 1. Регистрируем шрифт
        _register_fonts()
        
        # 2. Собираем данные
        data = _get_inspection_data(db, inspection_id)
        insp = data['inspection']
        
        status_map = {'draft': 'Черновик', 'submitted': 'Завершена', 'approved': 'Утверждена'}
        date_str = insp.date.strftime('%d.%m.%Y') if insp.date else '—'
        
        violations_html = ''
        violations_list = [v for v in data['violations'] if not v['is_safe']]
        all_safe = len(violations_list) == 0
        
        logger.info(f"📝 Нарушений: {len(violations_list)}, всё безопасно: {all_safe}")
        
        if all_safe:
            violations_html = '''
            <div style="text-align: center; padding: 30px; background-color: #d9f7be; border: 2px solid #52c41a;">
                <p style="font-size: 24px; color: #389e0d; font-weight: bold;">✓ ВСЁ БЕЗОПАСНО</p>
                <p style="color: #52c41a;">Нарушений не выявлено</p>
            </div>
            '''
        else:
            for v in violations_list:
                tags = ''
                if v['is_top']:
                    tags += '<font color="white" backColor="#cf1322"> ТОП-НАРУШЕНИЕ </font> '
                if v['is_gross']:
                    tags += '<font color="white" backColor="#fa8c16"> ГРУБЕЙШЕЕ </font> '
                if v['is_stopped']:
                    tags += '<font color="white" backColor="#cf1322"> ОСТАНОВКА РАБОТ </font> '
                
                photos_html = ''
                for photo in v['photos'][:4]:
                    photos_html += f'<img src="{photo["uri"]}" width="120" height="90" style="margin: 2px;" />'
                
                border_color = "#cf1322" if v["is_top"] else "#faad14"
                bg_color = "#fff1f0" if v["is_top"] else "#ffffff"
                
                violations_html += f'''
                <div style="border-left: 3px solid {border_color}; padding: 10px; margin-bottom: 15px; background-color: {bg_color};">
                    <p style="font-weight: bold; margin: 0 0 5px 0;">Нарушение #{v["order"]} {tags}</p>
                    <p style="margin: 3px 0;"><b>Описание:</b> {v["description"] or "—"}</p>
                    <p style="margin: 3px 0;"><b>Вид работ:</b> {v["work_type"] or "—"}</p>
                    {"<p style='margin: 3px 0;'><b>Нарушено ЗПБ:</b> " + v["zpb_rule"] + "</p>" if v["zpb_rule"] else ""}
                    {f'<div style="margin-top: 8px;">{photos_html}</div>' if photos_html else ""}
                </div>
                '''
        
        # Формируем абсолютный путь к шрифту для CSS
        font_uri = FONT_REGULAR_PATH.absolute().as_uri()
        
        html = f'''
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="utf-8">
            <style>
                @font-face {{
                    font-family: 'DejaVuSans';
                    src: url('{font_uri}');
                }}
                
                @page {{
                    size: A4;
                    margin: 2cm;
                }}
                body {{
                    font-family: 'DejaVuSans', sans-serif;
                    font-size: 11pt;
                    color: #333;
                }}
                h1 {{ color: #00A651; font-size: 20pt; margin-bottom: 5px; }}
                h2 {{ color: #003B5C; font-size: 14pt; border-bottom: 2px solid #00A651; padding-bottom: 3px; margin-top: 20px; }}
                .header {{ text-align: center; margin-bottom: 20px; border-bottom: 2px solid #00A651; padding-bottom: 10px; }}
                .company {{ font-size: 16pt; color: #00A651; font-weight: bold; }}
                .subtitle {{ font-size: 10pt; color: #666; }}
                table {{ width: 100%; border-collapse: collapse; margin-bottom: 15px; }}
                td, th {{ border: 1px solid #ccc; padding: 6px; font-size: 10pt; }}
                th {{ background-color: #f0f0f0; text-align: left; width: 35%; }}
                .footer {{ text-align: center; font-size: 9pt; color: #999; margin-top: 20px; border-top: 1px solid #ccc; padding-top: 5px; }}
            </style>
        </head>
        <body>
            <div class="header">
                <div class="company">ЕХ:Инспектра-ИПБ</div>
                <div class="subtitle">ГК «ЕВРОХИМ» · Система инспекций производственной безопасности</div>
            </div>
            <h1>Карта наблюдения №{insp.inspection_number}</h1>
            <table>
                <tr><th>Дата проведения</th><td>{date_str}</td></tr>
                <tr><th>Производственная единица</th><td>{data["pe"]}</td></tr>
                <tr><th>Подразделение</th><td>{data["department"]}</td></tr>
                <tr><th>Инспектор</th><td>{data["inspector"]}</td></tr>
                <tr><th>Подрядчик</th><td>{data["contractor"]}</td></tr>
                <tr><th>Место работ</th><td>{insp.work_location}</td></tr>
                <tr><th>Статус</th><td>{status_map.get(insp.status, insp.status)}</td></tr>
                <tr><th>Источник</th><td>{"Мобильное приложение" if insp.source == "mobile" else "Веб-интерфейс"}</td></tr>
            </table>
            <h2>Результаты проверки</h2>
            {violations_html}
            <div class="footer">
                Документ сформирован автоматически · {datetime.now().strftime("%d.%m.%Y %H:%M")} · ЕХ:Инспектра-ИПБ
            </div>
        </body>
        </html>
        '''
        
        logger.info(f"📄 HTML сгенерирован, размер: {len(html)} символов")
        logger.info(f"🔨 Запуск xhtml2pdf...")
        
        result = BytesIO()
        pisa_status = pisa.CreatePDF(
            html,
            dest=result,
            encoding='utf-8',
        )
        
        if pisa_status.err:
            logger.error(f"❌ Ошибка генерации PDF: {pisa_status.err}")
            raise RuntimeError(f"Ошибка генерации PDF: {pisa_status.err}")
        
        pdf_bytes = result.getvalue()
        logger.info(f"✅ PDF успешно сгенерирован: {len(pdf_bytes)} байт")
        
        return pdf_bytes
        
    except Exception as e:
        logger.error(f"❌ Критическая ошибка при генерации PDF: {e}", exc_info=True)
        raise
    finally:
        if data:
            _cleanup_temp_files(data.get('violations', []))

# ==================== DOCX EXPORT ====================
def generate_inspection_docx(inspection_id: uuid.UUID, db: Session) -> bytes:
    """Генерация DOCX-отчёта по карте наблюдения."""
    data = None
    try:
        logger.info(f"🚀 Начало генерации DOCX для проверки {inspection_id}")
        
        data = _get_inspection_data(db, inspection_id)
        insp = data['inspection']
        
        doc = Document()
        style = doc.styles['Normal']
        style.font.name = 'DejaVu Sans'
        style.font.size = Pt(11)
        
        header_para = doc.add_paragraph()
        header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = header_para.add_run('ГК «ЕВРОХИМ»')
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0, 166, 81)
        run.bold = True
        
        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitle.add_run('ЕХ:Инспектра-ИПБ · Карта наблюдения')
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(100, 100, 100)
        
        doc.add_paragraph()
        
        title = doc.add_heading(f'Карта наблюдения №{insp.inspection_number}', level=1)
        title.runs[0].font.color.rgb = RGBColor(0, 59, 92)
        
        table = doc.add_table(rows=8, cols=2)
        table.style = 'Light Grid Accent 1'
        
        status_map = {'draft': 'Черновик', 'submitted': 'Завершена', 'approved': 'Утверждена'}
        date_str = insp.date.strftime('%d.%m.%Y') if insp.date else '—'
        
        rows_data = [
            ('Дата проведения', date_str),
            ('Производственная единица', data['pe']),
            ('Подразделение', data['department']),
            ('Инспектор', data['inspector']),
            ('Подрядчик', data['contractor']),
            ('Место работ', insp.work_location),
            ('Статус', status_map.get(insp.status, insp.status)),
            ('Источник', 'Мобильное приложение' if insp.source == 'mobile' else 'Веб-интерфейс'),
        ]
        
        for i, (label, value) in enumerate(rows_data):
            table.rows[i].cells[0].text = label
            table.rows[i].cells[1].text = value
            for paragraph in table.rows[i].cells[0].paragraphs:
                for run in paragraph.runs:
                    run.bold = True
        
        doc.add_paragraph()
        
        doc.add_heading('Результаты проверки', level=2)
        
        violations_list = [v for v in data['violations'] if not v['is_safe']]
        
        if not violations_list:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run('✓ ВСЁ БЕЗОПАСНО')
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(82, 196, 26)
            run.bold = True
            
            p2 = doc.add_paragraph('Нарушений не выявлено. Инспектор подтвердил безопасные условия труда.')
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            stats = doc.add_paragraph()
            stats.add_run(f'Выявлено нарушений: ')
            run = stats.add_run(str(len(violations_list)))
            run.bold = True
            run.font.color.rgb = RGBColor(207, 19, 34)
            
            top_count = sum(1 for v in violations_list if v['is_top'])
            if top_count > 0:
                stats.add_run(' · ТОП-нарушений: ')
                run = stats.add_run(str(top_count))
                run.bold = True
                run.font.color.rgb = RGBColor(207, 19, 34)
            
            doc.add_paragraph()
            
            for v in violations_list:
                doc.add_heading(f'Нарушение #{v["order"]}', level=3)
                
                tags = []
                if v['is_top']: tags.append('ТОП-НАРУШЕНИЕ')
                if v['is_gross']: tags.append('ГРУБЕЙШЕЕ')
                if v['is_stopped']: tags.append('ОСТАНОВКА РАБОТ')
                
                if tags:
                    p = doc.add_paragraph()
                    run = p.add_run(' | '.join(tags))
                    run.font.color.rgb = RGBColor(207, 19, 34)
                    run.bold = True
                
                doc.add_paragraph(f'Описание: {v["description"] or "—"}')
                doc.add_paragraph(f'Вид работ: {v["work_type"] or "—"}')
                
                if v['zpb_rule']:
                    doc.add_paragraph(f'Нарушено ЗПБ: {v["zpb_rule"]}')
                
                if v['photos']:
                    doc.add_paragraph()
                    p = doc.add_paragraph()
                    run = p.add_run(f'Фотографии ({len(v["photos"])}):')
                    run.bold = True
                    
                    photo_cols = 2
                    photo_rows = (min(len(v['photos']), 6) + photo_cols - 1) // photo_cols
                    photo_table = doc.add_table(rows=photo_rows, cols=photo_cols)
                    photo_table.autofit = False
                    
                    for idx, photo in enumerate(v['photos'][:6]):
                        try:
                            row = idx // photo_cols
                            col = idx % photo_cols
                            cell = photo_table.rows[row].cells[col]
                            
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    run.text = ''
                            
                            paragraph = cell.paragraphs[0]
                            run = paragraph.add_run()
                            run.add_picture(photo['path'], width=Cm(5))
                            logger.info(f"✅ Фото вставлено в DOCX: {photo['path']}")
                            
                        except Exception as e:
                            logger.error(f"Ошибка вставки фото в DOCX: {e}", exc_info=True)
                            row = idx // photo_cols
                            col = idx % photo_cols
                            cell = photo_table.rows[row].cells[col]
                            cell.text = f"[Ошибка загрузки фото {idx + 1}]"
                
                doc.add_paragraph()
        
        doc.add_paragraph()
        
        footer = doc.add_paragraph()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run(f'Документ сформирован автоматически · {datetime.now().strftime("%d.%m.%Y %H:%M")} · ЕХ:Инспектра-ИПБ')
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(150, 150, 150)
        
        result = BytesIO()
        doc.save(result)
        docx_bytes = result.getvalue()
        
        logger.info(f"✅ DOCX успешно сгенерирован: {len(docx_bytes)} байт")
        
        return docx_bytes
        
    except Exception as e:
        logger.error(f"❌ Критическая ошибка при генерации DOCX: {e}", exc_info=True)
        raise
    finally:
        if data:
            _cleanup_temp_files(data.get('violations', []))